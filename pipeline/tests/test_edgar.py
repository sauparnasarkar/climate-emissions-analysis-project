import io
import json
import zipfile
from datetime import date

import numpy as np
import openpyxl
import pandas as pd
import pytest

from pipeline import edgar
from pipeline.common import RunReport
from pipeline.crosswalk import build_crosswalk, expanded_gaps

from .conftest import make_fetched

YEARS = [1988, 1989, 1990, 1991, 1992]
ENT = {"AAA": "Alpha", "BBB": "Beta", "AIR": "Int. Aviation", "SEA": "Int. Shipping"}


def workbook_zip(substance: str, values: dict, years, header_row: int = 10, extra_rows=()) -> bytes:
    """A TOTALS BY COUNTRY workbook shaped like EDGAR's: junk rows, header at `header_row`, then data.
    values: {(code, year): number or None}; one row per (code, substance)."""
    wb = openpyxl.Workbook()
    wb.active.title = "Citations and references"
    ws = wb.create_sheet("TOTALS BY COUNTRY")
    for r in range(1, header_row):
        ws.cell(r, 1, "Content:" if r == 1 else None)
    hdr = ["IPCC_annex", "C_group_IM24_sh", "Country_code_A3", "Name", "Substance"] + [f"Y_{y}" for y in years]
    for c, v in enumerate(hdr, 1):
        ws.cell(header_row, c, v)
    r = header_row + 1
    for code, name in ENT.items():
        row = ["Non-Annex_I", "x", code, name, substance] + [values.get((code, y)) for y in years]
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        r += 1
    for row in extra_rows:
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as z:
        z.writestr("data.xlsx", buf.getvalue())
        z.writestr("_readme.html", "<html/>")
    return zbuf.getvalue()


def make_values(total_extra_pre1990=10.0):
    co2 = {(c, y): 1000.0 for c in ENT for y in YEARS}
    ch4 = {(c, y): 10.0 for c in ENT for y in YEARS}
    n2o = {(c, y): 1.0 for c in ENT for y in YEARS}
    fg = {(c, y): (50.0 if y in (1990, 1991) else None) for c in ENT for y in YEARS}  # no F-gases 1988-89, nor 1992
    tot = {}
    for c in ENT:
        for y in YEARS:
            base = 1000 + 28 * 10 + 265 * 1
            tot[(c, y)] = base + (50 if y in (1990, 1991) else (total_extra_pre1990 if y < 1990 else 50))
    return co2, ch4, n2o, fg, tot


@pytest.fixture
def zips():
    co2, ch4, n2o, fg, tot = make_values()
    return {
        "co2": workbook_zip("CO2", co2, YEARS),
        "ch4": workbook_zip("CH4", ch4, YEARS),
        "n2o": workbook_zip("N2O", n2o, YEARS),
        "fgas": workbook_zip("GWP_100_AR5_HFC", fg, YEARS),
        "total": workbook_zip("GWP_100_AR5_GHG", tot, YEARS),
    }


@pytest.fixture
def wides(zips):
    return {k: edgar.read_totals(v)[0] for k, v in zips.items()}


ROOT_HTML = '<a href="EDGAR_2024_GHG/">x</a> <a href="EDGAR_2026_GHG/">y</a> <a href="EDGAR_2025_GHG/">z</a> <a href="v80_FT2022/">old</a>'
REL_HTML = (
    '<a href="IEA_EDGAR_CO2_1970_1992.zip">a</a><a href="IEA_EDGAR_CO2_m_1970_1992.zip">monthly-not-this</a>'
    '<a href="EDGAR_CH4_1970_1992.zip">b</a><a href="EDGAR_N2O_1970_1992.zip">c</a>'
    '<a href="EDGAR_AR5g_F-gases_1990_1992.zip">d</a><a href="EDGAR_F-gases_1990_1992.zip">raw-not-this</a>'
    '<a href="EDGAR_AR5_GHG_1970_1992.zip">e</a>'
)


# ---------------------------------------------------------------- discovery & parsing


def test_discover_release_picks_latest_year():
    assert edgar.discover_release(ROOT_HTML) == "EDGAR_2026_GHG"
    with pytest.raises(ValueError, match="no EDGAR_<year>_GHG"):
        edgar.discover_release('<a href="v80_FT2022/">')


def test_find_files_ignores_monthly_and_raw_fgas():
    f = edgar.find_files(REL_HTML)
    assert f == {
        "co2": "IEA_EDGAR_CO2_1970_1992.zip", "ch4": "EDGAR_CH4_1970_1992.zip", "n2o": "EDGAR_N2O_1970_1992.zip",
        "fgas": "EDGAR_AR5g_F-gases_1990_1992.zip", "total": "EDGAR_AR5_GHG_1970_1992.zip",
    }
    with pytest.raises(ValueError, match="no file matching"):
        edgar.find_files(REL_HTML.replace("EDGAR_N2O_1970_1992.zip", "other.zip"))


def test_read_totals_finds_header_wherever_it_is():
    co2, *_ = make_values()
    for header_row in (10, 4):
        wide, names = edgar.read_totals(workbook_zip("CO2", co2, YEARS, header_row=header_row))
        assert list(wide.columns) == YEARS and set(wide.index) == set(ENT) and names["AIR"] == "Int. Aviation"
        assert wide.loc["AAA", 1990] == 1000.0


def test_read_totals_sums_multiple_substance_rows():
    co2, *_ = make_values()
    extra = [["Non-Annex_I", "x", "AAA", "Alpha", "HFC-134a"] + [5.0] * len(YEARS)]
    wide, _ = edgar.read_totals(workbook_zip("HFC-125", co2, YEARS, extra_rows=extra))
    assert wide.loc["AAA", 1990] == 1005.0  # two substance rows for one entity are summed


def test_read_totals_missing_header_raises():
    wb = openpyxl.Workbook()
    wb.active.title = "TOTALS BY COUNTRY"
    wb.active.cell(1, 1, "nothing here")
    buf = io.BytesIO()
    wb.save(buf)
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w") as z:
        z.writestr("d.xlsx", buf.getvalue())
    with pytest.raises(ValueError, match="Country_code_A3"):
        edgar.read_totals(zb.getvalue())


# ---------------------------------------------------------------- tables & reconciliation


def test_build_tables_composition_and_fgas_gaps(wides):
    country, g = edgar.build_tables(wides, wides["total"])
    g = g.set_index("year")
    # 4 entities x 1000 Gg CO2 = 4 Mt; CH4 4 x 10 x 28 = 1.12 Mt; N2O 4 x 265 = 1.06 Mt; F-gas 4 x 50 = 0.2 Mt
    assert g.loc[1990, "co2_mt"] == pytest.approx(4.0)
    assert g.loc[1990, "ch4_mtco2e"] == pytest.approx(1.12) and g.loc[1990, "n2o_mtco2e"] == pytest.approx(1.06)
    assert g.loc[1990, "fgas_mtco2e"] == pytest.approx(0.2)
    assert g.loc[1990, "residual_mtco2e"] == pytest.approx(0.0, abs=1e-9)
    # F-gases are explicit nulls (not zero, not interpolated) where the file has no data
    assert g.loc[[1988, 1989, 1992], "fgas_mtco2e"].isna().all() and not g.loc[[1988, 1989, 1992], "fgas_available"].any()
    assert g.loc[[1988, 1992], "components_sum_mtco2e"].isna().all() and g.loc[[1988, 1992], "residual_pct"].isna().all()
    assert g.loc[1988, "total_ghg_mtco2e"] > 0  # total still comes from the combined workbook


def test_bunkers_excluded_from_national_total_but_in_global(wides):
    country, g = edgar.build_tables(wides, wides["total"])
    row = g[g["year"] == 1990].iloc[0]
    per_entity = 1595 / 1000
    assert row["total_ghg_mtco2e"] == pytest.approx(4 * per_entity)
    assert row["bunkers_mtco2e"] == pytest.approx(2 * per_entity)
    assert row["national_total_mtco2e"] == pytest.approx(2 * per_entity)
    assert set(country.loc[country["iso3"].isin(["AIR", "SEA"]), "entity_type"]) == {"bunker"}
    assert len(country) == len(ENT) * len(YEARS)


def test_build_tables_rejects_gas_file_missing_a_total_year(wides):
    wides["ch4"] = wides["ch4"].drop(columns=[1991])
    with pytest.raises(ValueError, match="ch4 file lacks year.*1991"):
        edgar.build_tables(wides, wides["total"])


def test_entity_absent_from_a_gas_file_counts_as_zero(wides):
    wides["n2o"] = wides["n2o"].drop(index=["BBB"])
    country, _ = edgar.build_tables(wides, wides["total"])
    row = country[(country["iso3"] == "BBB") & (country["year"] == 1990)].iloc[0]
    assert row["n2o_mtco2e"] == 0.0


def test_reconcile_notes_gaps_and_stays_quiet_within_tolerance(wides):
    _, g = edgar.build_tables(wides, wides["total"])
    rep = RunReport("edgar")
    edgar.reconcile(g, rep)
    assert rep.deviations == []
    assert any("no data for 1988-1989, 1992" in n for n in rep.notes)


def test_reconcile_flags_residual_beyond_tolerance(wides):
    wides["total"] = wides["total"] * 1.05  # combined total 5% above the per-gas sum
    _, g = edgar.build_tables(wides, wides["total"])
    rep = RunReport("edgar")
    edgar.reconcile(g, rep)
    assert len(rep.deviations) == 1 and "1990-1991" in rep.deviations[0] and "worst" in rep.deviations[0]


def test_ranges_helper():
    assert edgar._ranges([1970, 1971, 1972, 1975, 1977, 1978]) == "1970-1972, 1975, 1977-1978"


# ---------------------------------------------------------------- crosswalk


def owid_df():
    return pd.DataFrame(
        {
            "country": ["Alpha", "Beta Land", "Serbia", "Montenegro", "Monaco", "World", "International aviation", "International shipping", "Gamma"],
            "iso_code": ["AAA", "BBB", "SRB", "MNE", "MCO", np.nan, np.nan, np.nan, "GGG"],
        }
    )


def test_crosswalk_matches_bunkers_and_flags_many_to_one():
    names = {"AAA": "Alpha", "BBB": "Beta Land (EDGAR)", "SCG": "Serbia and Montenegro", "ZZZ": "Zeta Island", "AIR": "Int. Aviation", "SEA": "Int. Shipping"}
    cw = build_crosswalk(names, owid_df(), expanded=["Alpha", "Gamma"]).set_index("iso3")
    assert cw.loc["AAA", "match"] == "iso3_exact" and not cw.loc["AAA", "name_differs"]
    assert cw.loc["BBB", "match"] == "iso3_exact" and cw.loc["BBB", "name_differs"]
    assert cw.loc["AIR", "match"] == "bunker" and cw.loc["AIR", "owid_name"] == "International aviation"
    assert cw.loc["SEA", "owid_name"] == "International shipping"
    assert cw.loc["ZZZ", "match"] == "edgar_only" and cw.loc["ZZZ", "note"] == ""
    assert "Montenegro" in cw.loc["SCG", "note"] and "Serbia" in cw.loc["SCG", "note"]  # derived from names, not hardcoded
    assert set(cw[cw["match"] == "owid_only"].index) == {"SRB", "MNE", "MCO", "GGG"}
    assert cw.loc["AAA", "expanded"] and not cw.loc["BBB", "expanded"]
    assert "World" not in set(cw["owid_name"].dropna())  # OWID aggregates never enter the crosswalk


def test_expanded_gaps_lists_countries_without_exact_match():
    cw = build_crosswalk({"AAA": "Alpha"}, owid_df())
    assert expanded_gaps(cw, ["Alpha", "Gamma", "Nowhere"]) == ["Gamma", "Nowhere"]


# ---------------------------------------------------------------- end to end


def make_fetcher(zips):
    pages = {edgar.EDGAR_ROOT: ROOT_HTML, edgar.EDGAR_ROOT + "EDGAR_2026_GHG/": REL_HTML}
    base = edgar.EDGAR_ROOT + "EDGAR_2026_GHG/"
    files = {base + "IEA_EDGAR_CO2_1970_1992.zip": zips["co2"], base + "EDGAR_CH4_1970_1992.zip": zips["ch4"],
             base + "EDGAR_N2O_1970_1992.zip": zips["n2o"], base + "EDGAR_AR5g_F-gases_1990_1992.zip": zips["fgas"],
             base + "EDGAR_AR5_GHG_1970_1992.zip": zips["total"]}

    def f(url):
        if url in pages:
            return make_fetched(url, pages[url])
        fe = make_fetched(url, "")
        fe.content = files[url]
        return fe

    return f


def test_run_end_to_end(tmp_path, zips):
    rep = edgar.run(make_fetcher(zips), out_dir=str(tmp_path), provenance_path=str(tmp_path / "p.json"), today=date(1993, 6, 1), owid=owid_df(), expanded=["Alpha", "Gamma"])
    assert rep.records["edgar_country_annual"] == len(ENT) * len(YEARS) and rep.records["edgar_global_composition_annual"] == len(YEARS)
    glob = pd.read_csv(tmp_path / "edgar_global_composition_annual.csv")
    assert glob["year"].tolist() == YEARS and glob["fgas_mtco2e"].isna().sum() == 3
    assert (tmp_path / "country_crosswalk.csv").exists()
    assert any("expanded-set countries without an exact ISO3 match in EDGAR: Gamma" in d for d in rep.deviations)
    prov = json.loads((tmp_path / "p.json").read_text())
    p = prov["edgar_global_composition_annual"]
    assert p["coverage"] == [1988, 1992] and p["source_release"]["release"] == "EDGAR_2026_GHG"
    assert p["reconciliation"]["fgas_unavailable_years"] == [1988, 1989, 1992]
    assert "CC BY-NC-ND 4.0" in p["license"] and "IEA" in p["license"]  # the licence caution is carried in provenance
    assert prov["country_crosswalk"]["match_counts"]["bunker"] == 2


def test_run_flags_stale_release_and_missing_owid(tmp_path, zips, monkeypatch):
    monkeypatch.setattr(edgar, "_default_owid_and_expanded", lambda: (None, None))
    rep = edgar.run(make_fetcher(zips), out_dir=str(tmp_path), provenance_path=str(tmp_path / "p.json"), today=date(2030, 1, 1))
    assert any("behind 2030" in d for d in rep.deviations)
    assert any("crosswalk not built" in d for d in rep.deviations)
    assert not (tmp_path / "country_crosswalk.csv").exists()
